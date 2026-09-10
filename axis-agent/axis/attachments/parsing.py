"""Read Office/text files without executing macros, formulas, or embedded code.

Sections are lossless text slices with stable locations. Spreadsheet cells are
also retained as structured data; search text never becomes the source of truth.
"""
from __future__ import annotations

from datetime import date, datetime
import json
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET
import posixpath

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT = 8_000_000
MAX_SECTIONS = 20_000
SECTION_CHARS = 1200  # Embedding input is split again by the actual tokenizer.
READABLE = {".txt", ".md", ".csv", ".xlsx", ".docx", ".pptx", ".pdf"}
ACCEPTED = READABLE | {".doc", ".ppt", ".xls"}


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value if value is None or isinstance(value, (str, int, float, bool)) else str(value)


def workbook_layout(path):
    """Read merge coordinates without loading a workbook's cells into memory."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    relation = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    result = {}
    with zipfile.ZipFile(path) as archive:
        relationships = {node.attrib['Id']: node.attrib['Target'] for node in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for sheet in ET.fromstring(archive.read('xl/workbook.xml')).findall('s:sheets/s:sheet', ns):
            target = relationships[sheet.attrib[relation]]
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            merged = []
            with archive.open(target) as stream:
                for _, element in ET.iterparse(stream, events=('end',)):
                    if element.tag.endswith('}mergeCell'):
                        merged.append(element.attrib['ref'])
                    element.clear()
            result[sheet.attrib['name']] = dict(state=sheet.attrib.get('state', 'visible'), merged_ranges=merged)
    return result


def parse(path: Path) -> tuple[list[dict], list[str]]:
    suffix = path.suffix.lower()
    if suffix not in READABLE:
        return [], ["Legacy Office file: save as DOCX, PPTX, or XLSX to read its contents. The original remains available for website upload."]
    if suffix in {".docx", ".pptx", ".xlsx"}:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(x.file_size for x in entries) > 100 * 1024 * 1024:
                raise ValueError("Expanded Office file exceeds the 100 MiB / 10,000 entry limit.")
    sections, warnings, total = [], [], 0

    def add(location, text, *, sheet=None, row=None, cells=None):
        nonlocal total
        if not text.strip():
            return
        total += len(text)
        if total > MAX_TEXT:
            raise ValueError("Document exceeds the extracted text limit; split it into smaller files.")
        for start in range(0, len(text), SECTION_CHARS):
            if len(sections) >= MAX_SECTIONS:
                raise ValueError("Document has too many sections; split it into smaller files.")
            sections.append(dict(location=location + (f" [characters {start + 1}:{min(len(text), start + SECTION_CHARS)}]" if len(text) > SECTION_CHARS else ""),
                                 text=text[start:start + SECTION_CHARS], sheet=sheet, row_number=row, cells=cells if start == 0 else None))

    if suffix in {".txt", ".md"}:
        text = path.read_bytes().decode("utf-8-sig")
        heading = ""
        for index, line in enumerate(text.splitlines(), 1):
            if line.lstrip().startswith("#"):
                heading = line.strip()[:180]
            add(f"line {index}" + (f" | {heading}" if heading else ""), line)
    elif suffix == ".csv":
        import csv
        with path.open(encoding="utf-8-sig", newline="") as stream:
            headers = None
            for index, row in enumerate(csv.reader(stream), 1):
                if len(row) > 1000:
                    raise ValueError("CSV exceeds 1,000 columns.")
                headers = headers or row
                add(f"CSV row {index}", json.dumps(dict(zip([f"{i + 1}: {h}" for i, h in enumerate(headers)], row)), ensure_ascii=False),
                    sheet="CSV", row=index, cells=[dict(column=i + 1, value=v, display=v) for i, v in enumerate(row)])
    elif suffix == ".xlsx":
        import openpyxl
        layout = workbook_layout(path)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        values = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for ws in formulas:
                if layout[ws.title]['merged_ranges'] or layout[ws.title]['state'] != 'visible':
                    add(f'{ws.title}!layout', json.dumps(layout[ws.title]), sheet=ws.title)
                if (ws.max_row and ws.max_row > 50000 or ws.max_column and ws.max_column > 1000
                        or (ws.max_row or 0) * (ws.max_column or 0) > 2_000_000):
                    raise ValueError("Sheet exceeds 50,000 rows or 1,000 columns.")
                headers = ""
                for index, (row, cached) in enumerate(zip(ws.iter_rows(), values[ws.title].iter_rows()), 1):
                    cells = []
                    for col, (cell, cache) in enumerate(zip(row, cached), 1):
                        if cell.value is None:
                            continue
                        formula = cell.value if cell.data_type == "f" else None
                        value = scalar(cache.value if formula else cell.value)
                        display = str(value) if value is not None else ""
                        if isinstance(value, (int, float)) and cell.number_format and set(cell.number_format) == {"0"}:
                            display = str(int(value)).zfill(len(cell.number_format))
                        cells.append(dict(column=col, address=cell.coordinate, value=value, display=display,
                                          formula=formula, number_format=cell.number_format))
                    if not cells:
                        continue
                    body = json.dumps(cells, ensure_ascii=False)
                    if not headers:
                        headers = " | ".join(c["display"] for c in cells)[:400]
                    add(f"{ws.title}!row {index}", f"Sheet: {ws.title}; first populated row: {headers}\n{body}",
                        sheet=ws.title, row=index, cells=cells)
            warnings.append("Formula results are cached values, not recalculated. Check formulas, units, merged headers, and environment before applying values.")
        finally:
            formulas.close()
            values.close()
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        document = Document(path)
        heading = ""
        for index, block in enumerate(document.iter_inner_content(), 1):
            if isinstance(block, Table):
                headers = " | ".join(c.text for c in block.rows[0].cells)[:400] if block.rows else ''
                for number, row in enumerate(block.rows, 1):
                    add(f"block {index}, table row {number} | {heading}", f"First table row: {headers}\n" + " | ".join(c.text for c in row.cells))
            else:
                if block.style and block.style.name.startswith("Heading"):
                    heading = block.text[:180]
                add(f"block {index} | {heading}", block.text)
        warnings.append("Text and tables extracted. Images, embedded objects, and visual layout require visual review.")
    elif suffix == ".pptx":
        from pptx import Presentation
        def shapes(items, location):
            for shape in items:
                if hasattr(shape, "shapes"):
                    shapes(shape.shapes, location)
                if shape.has_text_frame:
                    add(location, shape.text)
                if shape.has_table:
                    headers = " | ".join(c.text for c in shape.table.rows[0].cells)[:400]
                    for index, row in enumerate(shape.table.rows, 1):
                        add(f"{location}, table row {index}", f"First table row: {headers}\n" + " | ".join(c.text for c in row.cells))
        for number, slide in enumerate(Presentation(path).slides, 1):
            shapes(slide.shapes, f"slide {number}")
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                add(f"slide {number}, notes", slide.notes_slide.notes_text_frame.text)
        warnings.append("Slide text, tables, and notes extracted. Diagrams and images require visual review.")
    else:
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDF: provide an unlocked copy for reading.")
        if len(reader.pages) > 2000:
            raise ValueError("PDF exceeds 2,000 pages.")
        for number, page in enumerate(reader.pages, 1):
            contents = page.get_contents()
            if contents and len(contents.get_data()) > 10 * 1024 * 1024:
                raise ValueError("PDF page content exceeds the extraction limit.")
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"Page {number} has no readable text; OCR is required.")
            add(f"page {number}", text)
        warnings.append("PDF table structure and image content are not guaranteed by text extraction.")
    return sections, warnings[:100]
