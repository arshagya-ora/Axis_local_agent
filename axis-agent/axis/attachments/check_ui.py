"""Opt-in attachment UI check using an isolated MV3 profile and fixture runner.

From axis-agent: uv run python -m axis.attachments.check_ui --executable PATH
"""
import argparse
from io import BytesIO
import json
from pathlib import Path
import secrets
import socket
import sys
import tempfile
import threading
import time

AGENT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(AGENT / 'tests'))


def fixtures():
    from openpyxl import Workbook
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfWriter
    files = []
    workbook = Workbook()
    workbook.active.append(['Setting', 'Value', 'Units'])
    workbook.active.append(['Timeout', 30, 'Minutes'])
    document = Document()
    document.add_paragraph('Configure the timeout to 30 minutes.')
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = 'Configuration guide'
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    for name, value in [('mapd.xlsx', workbook), ('guide.docx', document), ('slides.pptx', deck), ('scan.pdf', pdf)]:
        stream = BytesIO()
        (value.write if name.endswith('.pdf') else value.save)(stream)
        files.append(dict(name=name, mimeType='application/octet-stream', buffer=stream.getvalue()))
    files.extend(dict(name=name, mimeType='text/plain', buffer=text.encode()) for name, text in [
        ('workflow.md', '# Setup\nSet the timeout to 30 minutes.'), ('notes.txt', 'Supporting notes.'), ('config.csv', 'setting,value\ntimeout,30')])
    return files


def main():
    from playwright.sync_api import sync_playwright
    import uvicorn
    from axis.models import AxisConfig
    from axis.ui_runner import Runner
    from axis.ui_service import create_app, extension_id
    from axis.ui_store import Store
    from axis.attachments.search import Embeddings
    from test_ui_service import ControlledOrchestrator

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--executable', default=None)
    parser.add_argument('--output', type=Path, default=AGENT.parent / 'tmp' / 'attachment-feature-check')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    extension = AGENT.parent / 'browser-agent-bridge-main' / 'extension'
    identifier, token = extension_id(), secrets.token_urlsafe(32)
    with tempfile.TemporaryDirectory(prefix='attachment-ui-', dir=AGENT.parent / 'tmp') as directory:
        history = Store(Path(directory) / 'history.sqlite3')
        runner = Runner(history, AxisConfig.load(), factory=ControlledOrchestrator)
        listener = socket.socket()
        listener.bind(('127.0.0.1', 0))
        port = listener.getsockname()[1]
        app = create_app(history, runner, token=token, origin=f'chrome-extension://{identifier}', port=port)
        runner.attachments.embeddings = Embeddings(AGENT / '.axis-ui' / 'attachments' / 'model-cache')
        server = uvicorn.Server(uvicorn.Config(app, log_level='error', access_log=False))
        thread = threading.Thread(target=server.run, kwargs={'sockets': [listener]}, daemon=True)
        thread.start()
        deadline = time.monotonic() + 5
        while not server.started:
            if time.monotonic() > deadline:
                raise RuntimeError('Fixture server did not start')
            time.sleep(.02)
        report = dict(checks=[], page_errors=[])
        try:
            with sync_playwright() as playwright:
                context = playwright.chromium.launch_persistent_context(str(Path(directory) / 'profile'), headless=True,
                    executable_path=args.executable, args=[f'--disable-extensions-except={extension}', f'--load-extension={extension}'],
                    viewport={'width': 400, 'height': 900})
                try:
                    worker = context.service_workers[0] if context.service_workers else context.wait_for_event('serviceworker')
                    worker.evaluate('async ({url,token}) => chrome.storage.local.set({axisServiceUrl:url,axisUiToken:token})', {'url': f'http://127.0.0.1:{port}', 'token': token})
                    page = context.new_page()
                    page.on('pageerror', lambda error: report['page_errors'].append(str(error)))
                    page.goto(f'chrome-extension://{identifier}/sidepanel.html')
                    page.wait_for_function("() => document.querySelector('#connection-details').textContent.includes('Activity stream: live')")
                    page.locator('#attachment-composer input[type=file]').set_input_files(fixtures())
                    page.wait_for_function("() => document.querySelectorAll('.attachment-tile').length === 7 && !document.querySelector('#attachment-composer').textContent.includes('Reading file…')", timeout=30000)
                    page.get_by_label('Use workflow.md as', exact=True).select_option('instructions')
                    page.get_by_label('Use scan.pdf as', exact=True).select_option('upload')
                    assert page.locator('#attachment-composer').inner_text().count('Available for website upload') == 1
                    for width, height in [(400, 900), (320, 640)]:
                        page.set_viewport_size({'width': width, 'height': height})
                        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                        assert page.locator('#send').bounding_box()['y'] + page.locator('#send').bounding_box()['height'] <= height
                        page.screenshot(path=str(args.output / f'attachments-{width}.png'))
                    report['checks'].append('All seven formats attach; roles, unreadable-PDF notice, and narrow layout work')
                    page.get_by_label('Instruction', exact=True).fill('Follow workflow.md using MAPD values; upload scan.pdf at the documentation step.')
                    page.get_by_role('button', name='Send instruction', exact=True).click()
                    page.get_by_role('button', name='Pause', exact=True).wait_for()
                    assert page.locator('.attachment-tile').count() == 0
                    page.get_by_role('button', name='Pause', exact=True).click()
                    page.get_by_role('button', name='Resume', exact=True).wait_for()
                    page.reload()
                    page.get_by_role('button', name='Resume', exact=True).wait_for()
                    assert 'mapd.xlsx' in page.locator('.message.user').inner_text()
                    task = runner.active()
                    assert len(task['attachments']) == 7
                    assert any(x['filename'] == 'workflow.md' and x['role'] == 'instructions' for x in task['attachments'])
                    page.screenshot(path=str(args.output / 'attachment-history.png'))
                    report['checks'].append('Submission binds all attachments; pause and reload restore attachment history')
                    page.get_by_role('button', name='Stop', exact=True).click()
                    page.wait_for_function("() => ![...document.querySelectorAll('button')].some(x => x.textContent === 'Stop')")
                    assert not report['page_errors'], report['page_errors']
                finally:
                    context.close()
        finally:
            server.should_exit = True
            thread.join(timeout=30)
            (args.output / 'report.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
        print(json.dumps(report))


if __name__ == '__main__':
    main()
