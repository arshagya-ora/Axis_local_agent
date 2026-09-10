"""Attachment persistence using the UI's existing transactional SQLite store."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import hashlib
import heapq
import json
from pathlib import Path
import re
import shutil
import time
from uuid import uuid4

from axis.ui_store import ServiceError
from .parsing import ACCEPTED, parse
from .search import Embeddings, MODEL, fuse


class AttachmentStore:
    def __init__(self, history, directory: Path, *, embeddings=None):
        self.history, self.directory = history, Path(directory).resolve()
        self.directory.mkdir(parents=True, exist_ok=True)
        self.embeddings = embeddings or Embeddings(self.directory / "model-cache")
        self.worker = ThreadPoolExecutor(max_workers=1, thread_name_prefix="axis-attachments")
        self.indexer = ThreadPoolExecutor(max_workers=1, thread_name_prefix="axis-embeddings")
        self.jobs = set()
        with history.transaction():
            history.db.executescript("""
                CREATE TABLE IF NOT EXISTS attachments (
                    id TEXT PRIMARY KEY, conversation_id TEXT NOT NULL REFERENCES conversations(id) ON DELETE CASCADE,
                    filename TEXT NOT NULL, digest TEXT NOT NULL, size INTEGER NOT NULL,
                    status TEXT NOT NULL, search_status TEXT NOT NULL, warnings TEXT NOT NULL DEFAULT '[]', created_at REAL NOT NULL);
                CREATE TABLE IF NOT EXISTS attachment_sections (
                    id INTEGER PRIMARY KEY, attachment_id TEXT NOT NULL REFERENCES attachments(id) ON DELETE CASCADE,
                    ordinal INTEGER NOT NULL, location TEXT NOT NULL, text TEXT NOT NULL,
                    sheet TEXT, row_number INTEGER, cells TEXT, UNIQUE(attachment_id, ordinal));
                CREATE VIRTUAL TABLE IF NOT EXISTS attachment_fts USING fts5(text, location, content='attachment_sections', content_rowid='id');
                CREATE TRIGGER IF NOT EXISTS attachment_insert AFTER INSERT ON attachment_sections BEGIN
                    INSERT INTO attachment_fts(rowid,text,location) VALUES(new.id,new.text,new.location); END;
                CREATE TRIGGER IF NOT EXISTS attachment_delete AFTER DELETE ON attachment_sections BEGIN
                    INSERT INTO attachment_fts(attachment_fts,rowid,text,location) VALUES('delete',old.id,old.text,old.location); END;
                CREATE TABLE IF NOT EXISTS attachment_vectors (
                    section_id INTEGER PRIMARY KEY REFERENCES attachment_sections(id) ON DELETE CASCADE,
                    model TEXT NOT NULL, vector BLOB NOT NULL);
                CREATE TABLE IF NOT EXISTS embedding_cache (
                    digest TEXT NOT NULL, model TEXT NOT NULL, vector BLOB NOT NULL,
                    PRIMARY KEY(digest,model));
                CREATE TABLE IF NOT EXISTS task_attachments (
                    task_id TEXT NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
                    attachment_id TEXT NOT NULL REFERENCES attachments(id), role TEXT NOT NULL,
                    PRIMARY KEY(task_id, attachment_id));
                CREATE TABLE IF NOT EXISTS document_reads (
                    task_id TEXT NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
                    section_id INTEGER NOT NULL REFERENCES attachment_sections(id) ON DELETE CASCADE,
                    PRIMARY KEY(task_id,section_id));
                CREATE TABLE IF NOT EXISTS workflow_steps (
                    id TEXT PRIMARY KEY, task_id TEXT NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
                    section_id INTEGER NOT NULL REFERENCES attachment_sections(id),
                    ordinal INTEGER NOT NULL, instruction TEXT NOT NULL, expected_result TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending', evidence TEXT, goal_id TEXT,
                    verification TEXT,
                    UNIQUE(task_id,section_id,instruction));
            """)
            if 'verification' not in {row[1] for row in history.db.execute('PRAGMA table_info(workflow_steps)')}:
                history.db.execute('ALTER TABLE workflow_steps ADD COLUMN verification TEXT')
            # A stopped parser cannot leave attachments permanently spinning.
            history.db.execute("UPDATE attachments SET status='unreadable', search_status='keyword', warnings=? WHERE status='processing'",
                               (json.dumps(["Reading was interrupted. Remove and attach the file again; the original can still be uploaded."]),))
            history.db.execute("UPDATE attachments SET search_status='keyword' WHERE search_status='indexing'")
        self._collect_orphans()

    def _collect_orphans(self):
        with self.history.lock:
            known = {r[0] for r in self.history.db.execute("SELECT id FROM attachments")}
            for item in self.directory.iterdir():
                if re.fullmatch(r"[a-f0-9]{32}", item.name) and item.name not in known and item.is_dir() and not item.is_symlink():
                    self._remove_directory(item)

    def _remove_directory(self, path):
        resolved = path.resolve()
        if resolved.parent != self.directory or not re.fullmatch(r"[a-f0-9]{32}", resolved.name):
            raise ValueError("Attachment directory is outside the store.")
        shutil.rmtree(resolved, ignore_errors=True)

    def close(self):
        self.worker.shutdown(wait=True, cancel_futures=True)
        self.indexer.shutdown(wait=True, cancel_futures=True)

    def accept(self, conversation_id, filename, staged: Path, size):
        filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
        if not filename or len(filename) > 180 or any(ord(c) < 32 for c in filename) or Path(filename).suffix.lower() not in ACCEPTED:
            raise ServiceError("invalid_file", "Choose PDF, DOCX, PPTX, XLSX, MD, TXT, CSV, or a legacy Office file.", 422)
        # Store under a generated name; preserve a safe original basename for website uploads.
        filename = re.sub(r'[<>:"|?*]', "_", filename).rstrip(" .")
        if re.fullmatch(r"(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\..*)?", filename):
            filename = "file_" + filename
        digest = hashlib.sha256(staged.read_bytes()).hexdigest()
        identifier = uuid4().hex
        directory = self.directory / identifier
        with self.history.transaction():
            self.history.conversation(conversation_id)
            if self.history.db.execute("SELECT count(*) FROM attachments WHERE conversation_id=?", (conversation_id,)).fetchone()[0] >= 32:
                raise ServiceError("attachment_limit", "Remove an unused attachment before adding more (32 per conversation).", 422)
            directory.mkdir()
            try:
                staged.replace(directory / filename)
                self.history.db.execute("INSERT INTO attachments VALUES(?,?,?,?,?,'processing','indexing','[]',?)",
                                        (identifier, conversation_id, filename, digest, size, time.time()))
            except BaseException:
                self._remove_directory(directory)
                raise
        future = self.worker.submit(self._index, identifier)
        self.jobs.add(future)
        future.add_done_callback(self.jobs.discard)
        return self.get(identifier)

    def get(self, identifier):
        with self.history.lock:
            row = self.history.db.execute("SELECT a.*, (SELECT count(*) FROM attachment_sections s WHERE s.attachment_id=a.id) AS sections FROM attachments a WHERE id=?", (identifier,)).fetchone()
            if row is None:
                raise ServiceError("not_found", "Attachment no longer exists.", 404)
            item = dict(row)
            item["warnings"] = json.loads(item["warnings"])
            return item

    def path(self, identifier):
        item = self.get(identifier)
        path = (self.directory / identifier / item["filename"]).resolve()
        if path.parent != self.directory / identifier or not path.is_file():
            raise ServiceError("missing_file", "The original attachment is missing.", 404)
        return path

    def list(self, conversation_id):
        with self.history.lock:
            self.history.conversation(conversation_id)
            ids = [r[0] for r in self.history.db.execute("SELECT id FROM attachments WHERE conversation_id=? ORDER BY created_at", (conversation_id,))]
            return [self.get(identifier) for identifier in ids]

    def delete(self, identifier):
        with self.history.transaction():
            self.get(identifier)
            if self.history.db.execute("SELECT 1 FROM task_attachments WHERE attachment_id=?", (identifier,)).fetchone():
                raise ServiceError("attachment_in_use", "This file belongs to task history. Delete the conversation to remove it.")
            self.history.db.execute("DELETE FROM attachments WHERE id=?", (identifier,))
            self.history.db.execute('DELETE FROM embedding_cache')
        self._remove_directory(self.directory / identifier)

    def _index(self, identifier):
        try:
            sections, warnings = parse(self.path(identifier))
            with self.history.transaction():
                self.get(identifier)  # deletion may have raced the worker
                for ordinal, section in enumerate(sections, 1):
                    self.history.db.execute("INSERT INTO attachment_sections(attachment_id,ordinal,location,text,sheet,row_number,cells) VALUES(?,?,?,?,?,?,?)",
                                            (identifier, ordinal, section["location"], section["text"], section["sheet"], section["row_number"], json.dumps(section["cells"]) if section["cells"] is not None else None))
                self.history.db.execute("UPDATE attachments SET status=?,warnings=?,search_status=? WHERE id=?",
                                        ("ready" if sections else "unreadable", json.dumps(warnings), "indexing" if sections else "keyword", identifier))
            if sections:
                self.indexer.submit(self._embed, identifier)
        except Exception as error:
            # Never copy parser internals or document contents into public errors.
            warning = str(error) if isinstance(error, (ValueError, UnicodeError)) else f"Could not read this file ({type(error).__name__}). Provide an unlocked, supported document."
            with self.history.transaction():
                self.history.db.execute("UPDATE attachments SET status='unreadable',search_status='keyword',warnings=? WHERE id=?",
                                        (json.dumps([warning[:600]]), identifier))

    def _embed(self, identifier):
        try:
            with self.history.lock:
                rows = self.history.db.execute("SELECT id,text,location FROM attachment_sections WHERE attachment_id=? ORDER BY ordinal", (identifier,)).fetchall()
            for start in range(0, len(rows), 32):
                batch = rows[start:start + 32]
                texts = [r['location'] + '\n' + r['text'] for r in batch]
                digests = [hashlib.sha256(text.encode()).hexdigest() for text in texts]
                with self.history.lock:
                    cached = [self.history.db.execute('SELECT vector FROM embedding_cache WHERE digest=? AND model=?', (digest, MODEL)).fetchone() for digest in digests]
                missing = [i for i, value in enumerate(cached) if value is None]
                generated = self.embeddings.encode([texts[i] for i in missing]) if missing else []
                vectors = {i: vector.tobytes() for i, vector in zip(missing, generated)}
                with self.history.transaction():
                    self.get(identifier)
                    for i, row in enumerate(batch):
                        vector = cached[i][0] if cached[i] else vectors[i]
                        self.history.db.execute("INSERT OR REPLACE INTO attachment_vectors VALUES(?,?,?)", (row["id"], MODEL, vector))
                        self.history.db.execute('INSERT OR IGNORE INTO embedding_cache VALUES(?,?,?)', (digests[i], MODEL, vector))
            status = "hybrid"
        except Exception:
            status = "keyword"
        with self.history.transaction():
            self.history.db.execute("UPDATE attachments SET search_status=? WHERE id=?", (status, identifier))

    def validate_refs(self, conversation_id, refs, task_id=None):
        if len({ref.attachment_id for ref in refs}) != len(refs):
            raise ServiceError('invalid_attachment', 'Each attachment can appear once per message.', 422)
        if task_id:
            with self.history.lock:
                existing = {row[0] for row in self.history.db.execute('SELECT attachment_id FROM task_attachments WHERE task_id=?', (task_id,))}
            if len(existing | {r.attachment_id for r in refs}) > 8:
                raise ServiceError('attachment_limit', 'A task can use at most eight attachments.', 422)
        for ref in refs:
            item = self.get(ref.attachment_id)
            if item["conversation_id"] != conversation_id:
                raise ServiceError("invalid_attachment", "Attachments must belong to this conversation.", 422)
            if ref.role not in {"auto", "upload"} and item["status"] != "ready":
                raise ServiceError("attachment_not_ready", f"{item['filename']} is not ready for reading. Wait, or use it for website upload only.", 422)

    def link(self, task_id, conversation_id, refs):
        with self.history.transaction():
            self.validate_refs(conversation_id, refs, task_id)
            for ref in refs:
                existing = self.history.db.execute('SELECT role FROM task_attachments WHERE task_id=? AND attachment_id=?', (task_id, ref.attachment_id)).fetchone()
                role = 'upload' if existing and existing[0] == 'upload' and ref.role == 'auto' else ref.role
                self.history.db.execute("INSERT INTO task_attachments VALUES(?,?,?) ON CONFLICT(task_id,attachment_id) DO UPDATE SET role=excluded.role",
                                        (task_id, ref.attachment_id, role))

    def catalog(self, task_id):
        with self.history.lock:
            rows = self.history.db.execute("SELECT attachment_id,role FROM task_attachments WHERE task_id=?", (task_id,)).fetchall()
            return [{**self.get(row[0]), "role": row[1]} for row in rows]

    def require(self, task_id, identifier):
        with self.history.lock:
            row = self.history.db.execute("SELECT role FROM task_attachments WHERE task_id=? AND attachment_id=?", (task_id, identifier)).fetchone()
            if row is None:
                raise ValueError("Attachment is not available to this task.")
            return row[0]

    def sections(self, task_id, identifier, *, ids=(), offset=0, limit=8, sheet=None, cell_range=None):
        self.require(task_id, identifier)
        args = [identifier]
        where = "attachment_id=?"
        if ids:
            where += " AND id IN (" + ",".join("?" for _ in ids) + ")"
            args.extend(ids)
        if sheet is not None:
            where += " AND sheet=?"
            args.append(sheet)
        columns = None
        if cell_range:
            if not sheet:
                raise ValueError("A cell range requires a sheet name.")
            from openpyxl.utils.cell import range_boundaries
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
            if not all((min_col, min_row, max_col, max_row)) or max_row < min_row or max_col < min_col:
                raise ValueError("Use a finite cell range, such as A2:D12.")
            where += " AND row_number BETWEEN ? AND ?"
            args.extend([min_row, max_row])
            columns = (min_col, max_col)
        with self.history.lock:
            rows = self.history.db.execute(f"SELECT * FROM attachment_sections WHERE {where} ORDER BY ordinal LIMIT ? OFFSET ?", (*args, limit + 1, offset)).fetchall()
        result, used = [], 0
        for row in rows[:limit]:
            item = dict(row)
            item["cells"] = json.loads(item["cells"]) if item["cells"] else None
            if columns and item["cells"]:
                item["cells"] = [c for c in item["cells"] if columns[0] <= c["column"] <= columns[1]]
            # Avoid repeating full structured rows for every text slice.
            if item["cells"] and len(json.dumps(item["cells"])) > 8000:
                item["cells"] = None
                item["cells_note"] = "Read a narrower column range for structured cells; section text is intact."
            size = len(json.dumps(item))
            if result and used + size > 16000:
                break
            result.append(item)
            used += size
        return dict(items=result, has_more=len(rows) > len(result), next_offset=offset + len(result))

    def search(self, task_id, query, identifier=None):
        import numpy as np
        ids = [x["id"] for x in self.catalog(task_id) if x["role"] != "upload"]
        if identifier:
            if self.require(task_id, identifier) == 'upload':
                raise ValueError('This file is designated for website upload only.')
            ids = [identifier]
        if not ids:
            return dict(items=[], mode="keyword")
        terms = re.findall(r"[^\W_]+", query, re.UNICODE)[:24]
        match = " OR ".join('"' + word + '"' for word in terms)
        placeholders = ",".join("?" for _ in ids)
        with self.history.lock:
            lexical = self.history.db.execute(f"SELECT s.id FROM attachment_fts JOIN attachment_sections s ON s.id=attachment_fts.rowid WHERE attachment_fts MATCH ? AND s.attachment_id IN ({placeholders}) ORDER BY bm25(attachment_fts),s.id LIMIT 20", (match, *ids)).fetchall() if match else []
            vector_sql = f"FROM attachment_vectors v JOIN attachment_sections s ON s.id=v.section_id WHERE s.attachment_id IN ({placeholders}) AND v.model=?"
            vector_count = self.history.db.execute('SELECT count(*) ' + vector_sql, (*ids, MODEL)).fetchone()[0]
            total = self.history.db.execute(f"SELECT count(*) FROM attachment_sections WHERE attachment_id IN ({placeholders})", ids).fetchone()[0]
        semantic, mode = [], "keyword"
        if vector_count:
            try:
                query_vector = self.embeddings.encode([query], query=True)[0]
                best = []
                with self.history.lock:
                    cursor = self.history.db.execute('SELECT v.section_id,v.vector ' + vector_sql, (*ids, MODEL))
                    while batch := cursor.fetchmany(512):
                        matrix = np.stack([np.frombuffer(row[1], dtype=np.float32) for row in batch])
                        scores = matrix @ query_vector
                        best = heapq.nlargest(20, best + [(float(score), -row[0]) for score, row in zip(scores, batch)])
                semantic = [-key for _, key in best]
                mode = "hybrid" if vector_count == total else "hybrid_partial"
            except Exception:
                pass
        ranked = fuse([r[0] for r in lexical], semantic)
        with self.history.lock:
            items = [dict(self.history.db.execute("SELECT s.id,s.attachment_id,a.filename,s.location,s.text FROM attachment_sections s JOIN attachments a ON a.id=s.attachment_id WHERE s.id=?", (key,)).fetchone()) for key in ranked]
        return dict(items=items, mode=mode, note="Matches locate sources; read the section before using values. Ranking is not confidence.")
