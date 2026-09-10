"""Real local DOM + durable runner approvals; scripted model, no Gmail/provider access."""
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
import argparse
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import threading
import time

AGENT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(AGENT), str(AGENT / 'tests')]
from playwright.sync_api import sync_playwright
from pptx import Presentation
from starlette.testclient import TestClient
from axis.models import AxisConfig
from axis.ui_runner import Runner
from axis.ui_service import create_app
from axis.ui_store import Store
from test_axis_agent import Script, make, browse, done
from test_attachments import TinyEmbeddings, ingest
from test_ui_service import TOKEN, ORIGIN, HEADERS, conversation, submit

HTML = b'''<!doctype html><title>Mail fixture</title>
<button aria-label="Compose">Compose</button><section hidden>
<label>Message<textarea></textarea></label><label>Attach files<input type="file"></label>
<button aria-label="Send">Send</button></section>
<script>
const save = (operation, value) => fetch('/changes', {method:'POST', body:JSON.stringify({operation,value})});
document.querySelector('[aria-label=Compose]').onclick=async()=>{await save('create','draft');document.querySelector('section').hidden=false;};
document.querySelector('textarea').oninput=e=>save('autosave',e.target.value);
document.querySelector('input').onchange=async e=>{const f=e.target.files[0];const bytes=new Uint8Array(await crypto.subtle.digest('SHA-256',await f.arrayBuffer()));const sha256=Array.from(bytes,b=>b.toString(16).padStart(2,'0')).join('');await save('upload',{name:f.name,sha256});document.title='Draft saved with overview.pptx';};
document.querySelector('[aria-label=Send]').onclick=()=>save('send','forbidden');
</script>'''


class FixtureBridge:
    """Execute production browser-tool RPC shapes in an isolated local browser."""
    def __init__(self, executable, url):
        self.pool = ThreadPoolExecutor(max_workers=1)
        self.url, self.calls = url, []
        self.pool.submit(self.open, executable).result()

    def open(self, executable):
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(executable_path=executable, headless=True)
        self.page = self.browser.new_page()
        self.page.route('**/*', lambda route: route.continue_() if route.request.url.startswith(self.url) else route.abort())
        self.page.goto(self.url)

    def rpc(self, method, params=None, **_):
        return self.pool.submit(self.call, method, params or {}).result(15)

    def call(self, method, params):
        self.calls.append(method)
        if method in {'extension.info', 'native.sitePatterns'}:
            return {}
        if method == 'tabs.list':
            return {'tabs':[{'id':41, 'url':self.page.url, 'title':self.page.title(), 'active':True}]}
        if method == 'page.accessibilityTree':
            return {'snapshotId':'fixture', 'snapshot':self.page.locator('body').inner_text(), 'url':self.page.url, 'title':self.page.title()}
        if method == 'page.title':
            return {'title':self.page.title()}
        if method == 'expect.page.toHaveTitle':
            self.page.wait_for_function('(title)=>document.title===title', arg=params['title'])
            return {'passed':True, 'actual':self.page.title()}
        spec = params.get('locator') or {}
        if 'label' in spec:
            locator = self.page.get_by_label(spec['label'], exact=True)
        elif 'role' in spec:
            locator = self.page.get_by_role(spec['role'], name=spec.get('name'), exact=True)
        else:
            raise AssertionError((method, params))
        if method == 'locator.count':
            return {'count':locator.count(), 'elements':locator.evaluate_all('els=>els.map(e=>({tagName:e.tagName.toLowerCase(),type:e.type||"",ariaLabel:e.getAttribute("aria-label")||"",accessibleName:e.labels?.[0]?.textContent||e.textContent||""}))')}
        if method == 'locator.click':
            locator.click()
            self.page.locator('textarea').wait_for(state='visible')
        elif method == 'locator.fill':
            locator.fill(params['text'])
        elif method == 'locator.setInputFiles':
            locator.set_input_files(params['files'])
            self.page.wait_for_function('document.title==="Draft saved with overview.pptx"')
        else:
            raise AssertionError((method, params))
        return {}

    def close(self):
        def finish():
            self.browser.close()
            self.playwright.stop()
        self.pool.submit(finish).result()
        self.pool.shutdown()


def check(executable, mode):
    changes = []
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_): pass
        def do_GET(self):
            self.send_response(200); self.send_header('Content-Type','text/html'); self.end_headers(); self.wfile.write(HTML)
        def do_POST(self):
            changes.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
            self.send_response(200); self.end_headers(); self.wfile.write(b'{}')
    server = ThreadingHTTPServer(('127.0.0.1',0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    url = f'http://127.0.0.1:{server.server_port}'
    bridge = FixtureBridge(executable, url)
    try:
        with tempfile.TemporaryDirectory(prefix='axis-attachment-check-') as directory:
            store = Store(Path(directory)/'history.sqlite3')
            config = AxisConfig.load(); config.tools.approval_required_actions = []
            config.run.max_actions_per_step = 4
            scripts = []
            runner = Runner(store, config, factory=lambda cfg,on_event:make(scripts[0], config=cfg, bridge=bridge, on_event=on_event)[0])
            with TestClient(create_app(store, runner, token=TOKEN, origin=ORIGIN), base_url='http://127.0.0.1:8766', headers=HEADERS) as client:
                runner.attachments.embeddings = TinyEmbeddings()
                cid = conversation(client)
                summary = 'AXIS automates browser tasks and verifies results.'
                deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = summary
                original = BytesIO(); deck.save(original)
                item = ingest(runner.attachments,cid,'overview.pptx',original.getvalue())
                prompt = f'Prepare a draft at {url}/, summarize the attached file and attach its original. Do not send it.'
                target = lambda label: {'kind':'locator','locator':{'label':label}}
                calls = [('browser_act',{'tab':'tab_1','command':{'action':'click','target':target('Compose')}}),
                         ('browser_act',{'tab':'tab_1','command':{'action':'fill','target':target('Message'),'value':summary}}),
                         ('browser_act',{'tab':'tab_1','command':{'action':'upload','locator':{'label':'Attach files'},'files':[item['id']]}}),
                         ('browser_assert',{'tab':'tab_1','command':{'assertion':'title','expected':'Draft saved with overview.pptx'}})]
                script = Script([{'decision':'document','attachment_uses':[{'attachment_id':item['id'],'operations':['read','upload'],'user_evidence':prompt}],
                                  'document_request':{'operation':'read','attachment_id':item['id']}},
                                 browse('Prepare and verify draft',verification={'assertion':'title','expected':'Draft saved with overview.pptx'}),done('Verified draft with summary and original file.')],
                                [{'calls':calls,'outcome':{'status':'goal_reached','summary':'Draft saved and original attached'}}])
                scripts.append(script)
                response = submit(client,cid,text=prompt,execution_mode=mode,attachments=[{'attachment_id':item['id']}]); response.raise_for_status()
                task_id = response.json()['task']['id']; reviewed = set(); deadline = time.monotonic()+40
                while time.monotonic()<deadline:
                    task = client.get(f'/api/tasks/{task_id}').json()
                    pending = task.get('pending_approval')
                    if pending and pending['id'] not in reviewed:
                        assert mode=='approval', 'Automatic mode unexpectedly requested an execution approval'
                        before = len(changes); time.sleep(.15)
                        assert len(changes)==before==len(reviewed), 'Website changed before approval'
                        endpoint = f'/api/tasks/{task_id}/approvals/{pending["id"]}'
                        client.post(endpoint,json={'decision':'approve'}).raise_for_status()
                        assert client.post(endpoint,json={'decision':'approve'}).json()['duplicate']
                        reviewed.add(pending['id'])
                    if task['status'] in {'completed','failed','needs_user','paused','limit_reached'}: break
                    time.sleep(.02)
                assert task['status']=='completed', (task.get('current_activity'), runner.orchestrator.state.last_error)
                assert summary in script.planner_prompts[1], 'Summary must have document source evidence'
                assert [c['operation'] for c in changes]==['create','autosave','upload'], changes
                assert changes[1]['value']==summary and changes[2]['value']=={'name':'overview.pptx','sha256':hashlib.sha256(original.getvalue()).hexdigest()}
                assert len(reviewed)==(3 if mode=='approval' else 0)
                return {'mode':mode,'status':task['status'],'approvals':len(reviewed),'verified_changes':changes,'sent':False}
    finally:
        bridge.close(); server.shutdown(); server.server_close(); thread.join()


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--executable',required=True)
    args=parser.parse_args()
    print(json.dumps([check(args.executable,mode) for mode in ['approval','automatic']],indent=2))
