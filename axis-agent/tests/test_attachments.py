"""Attachment contracts: exact source reads, scoped retrieval, and durable progress."""
import json
from pathlib import Path
import sys
import time
from uuid import uuid4

import numpy as np
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from axis.attachments.models import DocumentRequest, WorkflowItem
from axis.attachments.parsing import parse
from axis.attachments.runtime import DocumentSession
from axis.attachments.search import fuse
from axis.attachments.store import AttachmentStore
from axis.models import AttachmentRef, AxisConfig, AssertionRecord, GoalCheck
from axis.ui_store import Store, ServiceError
from axis.orchestrator import AxisOrchestrator


class TinyEmbeddings:
    """Deterministic semantic synonym to prove independent vector retrieval."""
    def encode(self, texts, *, query=False):
        return np.asarray([[1, 0, 0] if any(word in text.lower() for word in ('logout', 'inactive')) else [0, 1, 0] for text in texts], dtype=np.float32)


class OfflineEmbeddings:
    def encode(self, *args, **kwargs):
        raise RuntimeError('offline')


@pytest.fixture
def documents(tmp_path):
    history = Store(tmp_path / 'history.sqlite3')
    files = AttachmentStore(history, tmp_path / 'attachments', embeddings=TinyEmbeddings())
    cid = history.create_conversation()['id']
    task = dict(id=uuid4().hex, conversation_id=cid, title='Fixture')
    with history.transaction():
        history.save_task(task)
    yield files, cid, task['id']
    files.close()
    history.close()


def ingest(files, cid, name, data):
    stage = files.directory / f'{uuid4().hex}.part'
    stage.write_bytes(data)
    item = files.accept(cid, name, stage, len(data))
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        item = files.get(item['id'])
        if item['status'] != 'processing' and item['search_status'] != 'indexing':
            return item
        time.sleep(.01)
    raise AssertionError('Indexing did not finish')


def link(files, cid, task, item, role='reference'):
    files.link(task, cid, [AttachmentRef(attachment_id=item['id'], role=role)])


def test_hybrid_search_finds_synonym_without_keyword_and_keeps_task_scope(documents):
    files, cid, task = documents
    item = ingest(files, cid, 'guide.md', b'Terminate inactive sessions after 30 minutes.\nUse blue for the theme.')
    link(files, cid, task, item)
    other = ingest(files, cid, 'unattached.txt', b'logout hidden configuration')
    result = files.search(task, 'logout')
    assert result['mode'] == 'hybrid'
    assert result['items'][0]['attachment_id'] == item['id']
    assert 'inactive' in result['items'][0]['text']
    assert all(x['attachment_id'] != other['id'] for x in result['items'])
    with pytest.raises(ValueError, match='not available'):
        files.search(task, 'anything', other['id'])


def test_offline_keyword_search_and_fts_query_escaping(documents):
    files, cid, task = documents
    files.embeddings = OfflineEmbeddings()
    item = ingest(files, cid, 'config.txt', b'SESSION_TIMEOUT is 30 minutes.')
    link(files, cid, task, item)
    assert item['search_status'] == 'keyword'
    result = files.search(task, '"SESSION_TIMEOUT" OR (missing:*)')
    assert result['mode'] == 'keyword'
    assert result['items'][0]['text'] == 'SESSION_TIMEOUT is 30 minutes.'
    assert files.search(task, '"()') == {'items': [], 'mode': 'keyword', 'note': 'Matches locate sources; read the section before using values. Ranking is not confidence.'}


def test_upload_original_bytes_and_no_arbitrary_paths(documents):
    files, cid, task = documents
    original = b'original\r\nbytes\r\n'
    item = ingest(files, cid, '../../config.txt', original)
    link(files, cid, task, item, 'upload')
    session = DocumentSession(files, task)
    path = Path(session.resolve_uploads([item['id']])[0])
    assert path.read_bytes() == original
    assert path.name == 'config.txt'
    with pytest.raises(ValueError):
        session.resolve_uploads([str(path)])
    with pytest.raises(ServiceError, match='task history'):
        files.delete(item['id'])


def test_cross_conversation_attachment_is_rejected(documents):
    files, cid, task = documents
    other = files.history.create_conversation()['id']
    item = ingest(files, other, 'other.txt', b'Other conversation')
    with pytest.raises(ServiceError, match='this conversation'):
        link(files, cid, task, item)


def test_read_pagination_covers_all_text_without_truncation(documents):
    files, cid, task = documents
    long_line = 'x' * 5000
    item = ingest(files, cid, 'long.txt', (long_line + '\n' + '\n'.join(f'step {i}' for i in range(15))).encode())
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    pieces, offset = [], 0
    while True:
        result = session.execute(DocumentRequest(operation='read', attachment_id=item['id'], offset=offset))
        pieces.extend(result['items'])
        if not result['has_more']:
            break
        offset = result['next_offset']
    assert ''.join(p['text'] for p in pieces if p['location'].startswith('line 1 ')) == long_line
    assert len(pieces) == item['sections']
    assert session.completion_errors()  # read != planned or executed


def test_workflow_requires_reading_order_and_verified_browser_postcondition(documents):
    files, cid, task = documents
    item = ingest(files, cid, 'steps.md', b'Set timeout to 30.\nSet theme to blue.')
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    rows = files.sections(task, item['id'])['items']
    request = DocumentRequest(operation='plan', steps=[WorkflowItem(section_id=r['id'], instruction=r['text'], expected_result='Value saved') for r in rows])
    with pytest.raises(ValueError, match='Read every'):
        session.execute(request)
    session.execute(DocumentRequest(operation='read', attachment_id=item['id']))
    session.execute(request)
    session.execute(request)  # retry cannot duplicate steps
    steps = session.workflow()['next_steps']
    assert len(steps) == 2
    with pytest.raises(ValueError, match='first pending'):
        session.begin(steps[1]['id'], 'goal_2')
    session.begin(steps[0]['id'], 'goal_1')
    assert session.context()['current_source']['items'][0]['text'] == 'Set timeout to 30.'
    state = AxisConfig.load().new_memory('Configure')
    state.current_goal_id = 'goal_1'
    state.current_verification = GoalCheck(assertion='value', expected='30', target='Timeout')
    with pytest.raises(ValueError, match='passing browser_assert'):
        session.verified(state, AxisOrchestrator._matches_goal_check)
    state.assertions.append(AssertionRecord(assertion='value', expected='30', target=json.dumps({'locator': {'label': 'Timeout'}}), passed=True, goal_id='goal_1'))
    session.verified(state, AxisOrchestrator._matches_goal_check)
    recovered = DocumentSession(files, task)
    assert recovered.workflow()['counts'] == {'pending': 1, 'verified': 1}
    assert recovered.workflow()['next_steps'][0]['id'] == steps[1]['id']
    assert recovered.completion_errors()


def test_spreadsheet_preserves_formula_cached_value_units_and_cell_locations(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / 'mapd.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Authentication'
    sheet.append(['Setting', 'Value', 'Unit'])
    sheet.append(['Idle timeout', 30, 'minutes'])
    sheet.append(['Identifier', 7, 'code'])
    sheet['B3'].number_format = '0000'
    sheet.append(['Calculated', '=B2*2', 'minutes'])
    workbook.save(path)
    rows, warnings = parse(path)
    assert rows[1]['cells'][1]['address'] == 'B2'
    assert rows[1]['cells'][1]['value'] == 30
    assert rows[2]['cells'][1]['display'] == '0007'
    assert rows[3]['cells'][1]['formula'] == '=B2*2'
    assert rows[3]['cells'][1]['value'] is None
    assert warnings


def test_word_powerpoint_pdf_csv_and_legacy_parsing(tmp_path):
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfWriter
    word = Document()
    word.add_heading('Authentication', 1)
    word.add_paragraph('Set session timeout to 30 minutes.')
    word.add_table(rows=1, cols=2).cell(0, 0).text = 'Configuration table'
    path = tmp_path / 'guide.docx'
    word.save(path)
    sections, _ = parse(path)
    assert any('30 minutes' in x['text'] for x in sections)
    assert any('table row' in x['location'] for x in sections)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Install the product'
    slide.notes_slide.notes_text_frame.text = 'Then configure authentication'
    path = tmp_path / 'guide.pptx'
    deck.save(path)
    sections, _ = parse(path)
    assert any(x['location'] == 'slide 1' for x in sections)
    assert any('configure authentication' in x['text'] for x in sections)
    path = tmp_path / 'scan.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    sections, warnings = parse(path)
    assert not sections and any('OCR' in text for text in warnings)
    path = tmp_path / 'config.csv'
    path.write_text('setting,value\ntimeout,30\n', encoding='utf-8')
    sections, _ = parse(path)
    assert sections[1]['cells'][1]['value'] == '30'
    assert parse(tmp_path / 'legacy.doc')[1]


def test_corrupt_file_stays_available_for_website_upload(documents):
    files, cid, task = documents
    item = ingest(files, cid, 'broken.pdf', b'not actually a PDF')
    assert item['status'] == 'unreadable'
    assert files.path(item['id']).read_bytes() == b'not actually a PDF'
    with pytest.raises(ServiceError, match='not ready'):
        link(files, cid, task, item)
    link(files, cid, task, item, 'upload')


def test_fusion_deduplicates_and_combines_independent_rankings():
    assert fuse([1, 2, 2], [2, 3])[0] == 2
    assert len(fuse([1, 2], [2, 3])) == 3


def test_attachment_api_streams_over_chat_limit_and_preserves_retry_identity(tmp_path):
    from starlette.testclient import TestClient
    from axis.ui_service import create_app
    from axis.ui_runner import Runner
    from test_ui_service import ControlledOrchestrator, TOKEN, ORIGIN, HEADERS, conversation, submit
    history = Store(tmp_path / 'history.sqlite3')
    runner = Runner(history, AxisConfig.load(), factory=ControlledOrchestrator)
    app = create_app(history, runner, token=TOKEN, origin=ORIGIN)
    runner.attachments.embeddings = OfflineEmbeddings()
    with TestClient(app, base_url='http://127.0.0.1:8766', headers=HEADERS) as client:
        cid = conversation(client)
        data = b'Long instruction.\n' * 3000
        response = client.post(f'/api/conversations/{cid}/attachments?filename=steps.txt', content=data)
        assert response.status_code == 201
        identifier = response.json()['id']
        refs = [dict(attachment_id=identifier, role='upload')]
        request_id = uuid4().hex
        accepted = submit(client, cid, attachments=refs, client_request_id=request_id)
        assert accepted.status_code == 202
        duplicate = submit(client, cid, attachments=refs, client_request_id=request_id)
        assert duplicate.json()['duplicate'] is True
        changed = submit(client, cid, client_request_id=request_id)
        assert changed.status_code == 409
        snapshot = client.get(f'/api/conversations/{cid}').json()
        assert snapshot['messages'][0]['attachments'] == refs
        assert snapshot['tasks'][0]['attachments'][0]['filename'] == 'steps.txt'
        assert client.get(f'/api/attachments/{identifier}', headers={'Authorization': 'Bearer wrong'}).status_code == 401
        client.post('/api/tasks/' + accepted.json()['task']['id'] + '/control', json={'action': 'stop'})


@pytest.mark.anyio
async def test_document_only_agent_reads_attachment_without_browser(documents):
    from test_axis_agent import Script, make, done
    files, cid, task = documents
    item = ingest(files, cid, 'guide.txt', b'The required timeout is 30 minutes.')
    link(files, cid, task, item)
    script = Script([dict(decision='document', document_request=dict(operation='read', attachment_id=item['id'])), done('The timeout is 30 minutes.')])
    orchestrator, bridge, _ = make(script)
    orchestrator.documents = DocumentSession(files, task)
    result = await orchestrator.start_task('What is the timeout in the attachment?', task_id=task)
    assert result.status == 'completed'
    assert bridge.calls == []
    assert '30 minutes' in script.planner_prompts[1]
    assert any(e.kind == 'document' and e.attachment_id == item['id'] for e in orchestrator.state.evidence)


@pytest.fixture
def anyio_backend():
    return 'asyncio'


@pytest.mark.anyio
async def test_planned_browser_step_only_completes_after_exact_assertion(documents):
    from test_axis_agent import Script, make, browse, done
    files, cid, task = documents
    item = ingest(files, cid, 'steps.txt', b'Confirm the page title is Google.')
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    section = session.execute(DocumentRequest(operation='read', attachment_id=item['id']))['items'][0]
    session.execute(DocumentRequest(operation='plan', steps=[WorkflowItem(section_id=section['id'], instruction=section['text'], expected_result='Title is Google')]))
    step_id = session.workflow()['next_steps'][0]['id']
    script = Script([browse('Confirm page title', workflow_step_id=step_id, verification={'assertion': 'title', 'expected': 'Google'}), done('Verified.')],
                    [{'calls': [('browser_assert', {'tab': 'tab_1', 'command': {'assertion': 'title', 'expected': 'Google'}})], 'outcome': {'status': 'goal_reached', 'summary': 'Title verified'}}])
    orchestrator, _, _ = make(script)
    orchestrator.documents = session
    result = await orchestrator.start_task('Follow the attached steps.', task_id=task)
    assert result.status == 'completed'
    assert session.workflow()['counts'] == {'verified': 1}
    assert not session.completion_errors()


@pytest.mark.anyio
async def test_browser_upload_resolves_attachment_id_to_original_file(documents):
    from test_axis_agent import Script, make, browse, done
    files, cid, task = documents
    item = ingest(files, cid, 'config.txt', b'Preserve this original file.')
    link(files, cid, task, item, 'upload')
    upload = ('browser_act', {'tab': 'tab_1', 'command': {'action': 'upload', 'locator': {'selector': 'input[type=file]'}, 'files': [item['id']]}})
    assertion = ('browser_assert', {'tab': 'tab_1', 'command': {'assertion': 'title', 'expected': 'Google'}})
    script = Script([browse('Upload the attached file', verification={'assertion': 'title', 'expected': 'Google'}), done('Uploaded.')],
                    [{'calls': [upload, assertion], 'outcome': {'status': 'goal_reached', 'summary': 'Receipt checked'}}])
    orchestrator, bridge, _ = make(script)
    orchestrator.documents = DocumentSession(files, task)
    orchestrator.browser.upload_roots = (files.path(item['id']).parent,)
    result = await orchestrator.start_task('Upload the attached file.', task_id=task)
    assert result.status == 'completed'
    dispatched = [params for method, params in bridge.calls if method == 'locator.setInputFiles']
    assert len(dispatched) == 1
    assert Path(dispatched[0]['files'][0]).read_bytes() == b'Preserve this original file.'


def test_recovery_cannot_weaken_saved_postcondition(documents):
    files, cid, task = documents
    item = ingest(files, cid, 'steps.txt', b'Set timeout to 30 minutes.')
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    row = session.execute(DocumentRequest(operation='read', attachment_id=item['id']))['items'][0]
    session.execute(DocumentRequest(operation='plan', steps=[WorkflowItem(section_id=row['id'], instruction=row['text'], expected_result='30 minutes saved')]))
    step = session.workflow()['next_steps'][0]
    check = GoalCheck(assertion='value', expected='30', target='Timeout')
    session.begin(step['id'], 'goal_1', check)
    recovered = DocumentSession(files, task)
    assert recovered.workflow()['next_steps'][0]['status'] == 'running'
    with pytest.raises(ValueError, match='Retain the saved'):
        recovered.begin(step['id'], 'goal_2', GoalCheck(assertion='title', expected='Settings'))
    recovered.begin(step['id'], 'goal_2', check)


def test_original_files_removed_when_conversation_is_deleted(tmp_path):
    from starlette.testclient import TestClient
    from axis.ui_service import create_app
    from axis.ui_runner import Runner
    from test_ui_service import ControlledOrchestrator, TOKEN, ORIGIN, HEADERS, conversation
    history = Store(tmp_path / 'history.sqlite3')
    runner = Runner(history, AxisConfig.load(), factory=ControlledOrchestrator)
    app = create_app(history, runner, token=TOKEN, origin=ORIGIN)
    runner.attachments.embeddings = OfflineEmbeddings()
    with TestClient(app, base_url='http://127.0.0.1:8766', headers=HEADERS) as client:
        cid = conversation(client)
        file = ingest(runner.attachments, cid, 'private.txt', b'Private fixture data')
        path = runner.attachments.path(file['id'])
        assert client.delete(f'/api/conversations/{cid}').status_code == 200
        assert not path.exists()
        assert client.get('/api/attachments/' + file['id']).status_code == 404


def test_index_reuses_embedding_cache(documents):
    files, cid, task = documents
    first = ingest(files, cid, 'one.txt', b'Cached text')
    assert first['search_status'] == 'hybrid'
    files.embeddings = OfflineEmbeddings()
    second = ingest(files, cid, 'two.txt', b'Cached text')
    assert second['search_status'] == 'hybrid'


def test_restart_recovers_running_step_without_replaying_verified_step(tmp_path):
    from starlette.testclient import TestClient
    from axis.ui_service import create_app
    from axis.ui_runner import Runner
    from test_ui_service import TOKEN, ORIGIN, HEADERS, wait_for
    from test_axis_agent import Script, make, browse, done
    path = tmp_path / 'history.sqlite3'
    history = Store(path)
    files = AttachmentStore(history, tmp_path / 'attachments', embeddings=TinyEmbeddings())
    cid, task_id = history.create_conversation()['id'], uuid4().hex
    config = AxisConfig.load()
    task = dict(id=task_id, conversation_id=cid, title='Configuration', instruction='Follow all steps in the attachment.',
                status='active', owns_runtime=True, created_at=time.time(), updated_at=time.time(),
                limits=config.run.model_dump(), counters={'browser_actions': 4, 'total_steps': 2, 'model_requests': 5})
    with history.transaction():
        history.save_task(task)
        history.message(cid, task_id, task['instruction'], uuid4().hex, 'new_task')
    item = ingest(files, cid, 'steps.txt', b'Confirm first setting.\nConfirm page title is Google.')
    link(files, cid, task_id, item, 'instructions')
    session = DocumentSession(files, task_id)
    rows = session.execute(DocumentRequest(operation='read', attachment_id=item['id']))['items']
    session.execute(DocumentRequest(operation='plan', steps=[WorkflowItem(section_id=row['id'], instruction=row['text'], expected_result='Configured') for row in rows]))
    first, second = session.workflow()['next_steps']
    with history.transaction():
        history.db.execute("UPDATE workflow_steps SET status='verified' WHERE id=?", (first['id'],))
    check = GoalCheck(assertion='title', expected='Google')
    session.begin(second['id'], 'old_goal', check)
    task.update(attachments=files.catalog(task_id), workflow=session.workflow())
    with history.transaction():
        history.save_task(task)
    files.close()
    history.close()

    history = Store(path)
    history.interrupt()
    script = Script([browse('Reconcile page title', workflow_step_id=second['id'], verification=check.model_dump()), done('Recovered and verified.')],
                    [{'calls': [('browser_assert', {'tab': 'tab_1', 'command': {'assertion': 'title', 'expected': 'Google'}})], 'outcome': {'status': 'goal_reached', 'summary': 'Existing result verified'}}])
    orchestrator, bridge, _ = make(script)
    def factory(config, on_event):
        orchestrator.on_event = on_event
        return orchestrator
    runner = Runner(history, config, factory=factory)
    with TestClient(create_app(history, runner, token=TOKEN, origin=ORIGIN), base_url='http://127.0.0.1:8766', headers=HEADERS) as client:
        response = client.post(f'/api/tasks/{task_id}/control', json={'action': 'resume'})
        assert response.status_code == 200
        wait_for(lambda: runner.future.done())
        saved = client.get(f'/api/tasks/{task_id}').json()
        assert saved['status'] == 'completed'
        assert saved['workflow']['counts'] == {'verified': 2}
        assert saved['counters']['browser_actions'] == 5
        assert saved['counters']['total_steps'] == 3
        assert saved['counters']['model_requests'] > 5
        assert 'locator.fillRef' not in bridge.methods


def test_upload_size_limit_cleans_staging_files(tmp_path, monkeypatch):
    from starlette.testclient import TestClient
    from axis.ui_service import create_app
    from axis.ui_runner import Runner
    from test_ui_service import ControlledOrchestrator, TOKEN, ORIGIN, HEADERS, conversation
    monkeypatch.setattr('axis.attachments.api.MAX_FILE_BYTES', 100)
    history = Store(tmp_path / 'history.sqlite3')
    runner = Runner(history, AxisConfig.load(), factory=ControlledOrchestrator)
    with TestClient(create_app(history, runner, token=TOKEN, origin=ORIGIN), base_url='http://127.0.0.1:8766', headers=HEADERS) as client:
        cid = conversation(client)
        response = client.post(f'/api/conversations/{cid}/attachments?filename=large.txt', content=b'a' * 101)
        assert response.status_code == 413
        assert not list(runner.attachments.directory.glob('*.part'))
        assert client.get(f'/api/conversations/{cid}/attachments').json()['items'] == []


def test_invalid_followup_attachments_are_rejected_before_acceptance(documents):
    files, cid, task = documents
    items = [ingest(files, cid, f'{i}.txt', b'Configuration') for i in range(9)]
    for item in items[:8]:
        link(files, cid, task, item)
    with pytest.raises(ServiceError, match='eight'):
        files.validate_refs(cid, [AttachmentRef(attachment_id=items[8]['id'])], task)
    duplicate = AttachmentRef(attachment_id=items[0]['id'])
    with pytest.raises(ServiceError, match='once'):
        files.validate_refs(cid, [duplicate, duplicate], task)
    assert len(files.catalog(task)) == 8


def test_merged_workbook_headers_retain_anchor_coordinates(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells('A1:C1')
    sheet['A1'] = 'Authentication module'
    sheet.append(['Setting', 'Value', 'Unit'])
    sheet.append(['Timeout', 30, 'minutes'])
    path = tmp_path / 'merged.xlsx'
    workbook.save(path)
    rows, _ = parse(path)
    metadata = json.loads(rows[0]['text'])
    assert metadata['merged_ranges'] == ['A1:C1']
    assert any('Authentication module' in row['text'] for row in rows)
