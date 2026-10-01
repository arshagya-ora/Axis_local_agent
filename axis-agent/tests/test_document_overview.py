"""Scale regression for the reported 915-section DOCX orientation request."""
import json
import sys
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from axis.attachments.models import DocumentRequest
from axis.attachments.runtime import DocumentSession, DocumentRestriction
from axis.models import AttachmentUse
from test_attachments import documents, ingest, link
from test_attachment_workflows import instruction, followup
from test_axis_agent import Script, make, done


@pytest.fixture
def anyio_backend():
    return 'asyncio'


def long_document(files, cid, task):
    from docx import Document
    doc = Document()
    for i in range(915):
        if i % 30 == 0:
            doc.add_heading(f'Topic {i//30}', level=1)
        else:
            doc.add_paragraph(f'Fact {i}. ' + 'Architecture, runtime, and delivery details. '*12)
    stream = BytesIO()
    doc.save(stream)
    item = ingest(files, cid, 'master-plan.docx', stream.getvalue())
    link(files, cid, task, item)
    return item


@pytest.mark.anyio
async def test_reported_prompt_uses_one_bounded_overview_instead_of_915_section_traversal(documents):
    files, cid, task = documents
    item = long_document(files, cid, task)
    prompt = 'what is written in this file?'
    instruction(files, task, prompt)
    session = DocumentSession(files, task)
    # Reproduce the old live planner's overly exhaustive first declaration.
    script = Script([dict(decision='document', attachment_uses=[dict(attachment_id=item['id'], operations=['read'],
        user_evidence=prompt, purpose='summary', coverage='all')],
        document_request=dict(operation='outline', attachment_id=item['id'])), done('An overview based on sampled text.')])
    orchestrator, bridge, _ = make(script)
    orchestrator.documents = session
    result = await orchestrator.start_task(prompt, task_id=task)
    assert result.status == 'completed' and result.model_requests == 2 and result.total_steps == 1
    assert not bridge.calls
    preview = session.saved_overview(item['id'])
    assert preview['total_sections'] == 915 and not preview['exhaustive']
    assert preview['sampled_sections'] <= 36
    assert preview['structure'][0]['label'] == 'Topic 0'
    assert preview['structure'][-1]['label'] == 'Topic 30'
    assert 'Fact 914.' in json.dumps(preview)
    assert len(json.dumps(preview)) < 32000
    assert not session.coverage()[0]['complete'] and session.coverage()[0]['read'] == 0
    assert not session.completion_errors()
    recovered = DocumentSession(files, task)
    assert recovered.saved_overview(item['id']) == preview
    # Do not send the same source packet in both last_result and overviews.
    assert 'excerpts' not in recovered.context()['last_result']


@pytest.mark.parametrize('prompt', ['List every scenario in this file.', 'Read every section of this file.', 'Provide a detailed summary of this file.'])
def test_overview_never_satisfies_exhaustive_requests(documents, prompt):
    files, cid, task = documents
    item = ingest(files, cid, 'scenarios.txt', '\n'.join(f'Scenario {i}' for i in range(35)).encode())
    link(files, cid, task, item)
    instruction(files, task, prompt)
    session = DocumentSession(files, task)
    session.execute(DocumentRequest(operation='overview', attachment_id=item['id']))
    assert session.completion_errors()
    assert session._uses()[item['id']]['coverage'] == 'all'
    assert session.coverage()[0]['read'] == 0


def test_preview_obeys_prompt_restrictions_after_followup(documents):
    files, cid, task = documents
    item = ingest(files, cid, 'source.txt', b'Private source content')
    link(files, cid, task, item)
    instruction(files, task, 'Give a brief overview of this file.')
    session = DocumentSession(files, task)
    session.execute(DocumentRequest(operation='overview', attachment_id=item['id']))
    followup(files, cid, task, 'Do not read this file.')
    session.refresh_scope()
    assert not session.context()['overviews'] and session.last_result is None
    with pytest.raises(DocumentRestriction):
        session.execute(DocumentRequest(operation='overview', attachment_id=item['id']))
